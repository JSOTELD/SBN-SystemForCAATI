"""Guías diarias y detalle individual para auditoría técnica.

Fuente: archivo de guías de 30 días, seis hojas GO, filas 13:42.
El usuario declara PRETEST real y la fuente identifica ambas fases como no verificadas:
se conserva la discrepancia para revisión documental.
No se transforman agregados en observaciones reales ni se modifica el censo.
PRCC es una condición conjunta; no permite inferir completitud y consistencia
por separado. Las asociaciones individuales se conservan como detalle técnico.
"""
import hashlib
import io
import json
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook
from ..extensions import db
from ..models import Asset, AuditLog, ResearchSample, GuideStudy, GuideAsset

ALGORITHM = 'daily-derivation-v1'
SEED = 'GUIAS-N1693-20260913-v1'
METRICS = ['PRCC', 'TPGR', 'PACI', 'TPI', 'PACR', 'PRA']
SHEETS = {'PRCC':'GO-01_PRCC','TPGR':'GO-02_TPGR','PACI':'GO-03_PACI','TPI':'GO-04_TPI','PACR':'GO-05_PACR','PRA':'GO-06_PRA'}
PERCENTAGES = {'PRCC', 'PACI', 'PACR', 'PRA'}


def canonical(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':')).encode('utf-8')


def digest(value):
    return hashlib.sha256(canonical(value)).hexdigest()


def cell_value(value):
    if isinstance(value, (datetime, date)): return value.isoformat()
    return value


def milliseconds(value):
    return int((Decimal(str(value)) * 1000).quantize(Decimal('1'), rounding=ROUND_HALF_UP))


def integer(value, label):
    if isinstance(value, bool) or not isinstance(value, (float, int)) or value != int(value) or value < 0:
        raise ValueError(f'{label}: se esperaba un entero no negativo.')
    return int(value)


def parse_guides(blob):
    workbook = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    formulas = load_workbook(io.BytesIO(blob), read_only=True, data_only=False)
    guides = {}; archive = {}; discrepancies = []
    try:
        for metric in METRICS:
            sheet = SHEETS[metric]
            values = list(workbook[sheet].values); formula_rows = list(formulas[sheet].values)
            if values[8][1] != 1693: raise ValueError(f'{sheet}: población diferente de 1693.')
            archive[sheet] = {'values': [[cell_value(v) for v in row] for row in values],
                              'formulas': [[cell_value(v) for v in row] for row in formula_rows]}
            rows = []
            for day, row in enumerate(values[12:42], 1):
                pair = {'day': day, 'sourceRow': day + 12}
                for phase, offset in [('PRETEST',0),('POSTTEST',6 if metric == 'TPGR' else 5)]:
                    stamp = row[offset]
                    if not isinstance(stamp, datetime): raise ValueError(f'{sheet}: falta fecha del día {day}.')
                    if metric == 'TPGR':
                        denominator = integer(row[offset+3], f'{sheet}/NR/{day}')
                        numerator = float(row[offset+4]); cached = row[offset+5]
                        code = row[offset+1]; profile = row[offset+2]
                        if code != f'TPGR-{day:02d}': raise ValueError(f'{sheet}: código de par no esperado.')
                        source_range = f'{"A" if offset == 0 else "G"}{day+12}:{"F" if offset == 0 else "L"}{day+12}'
                    else:
                        denominator = integer(row[offset+1], f'{sheet}/total/{day}')
                        numerator = float(row[offset+2]); cached = row[offset+3]
                        if row[offset+4] != day: raise ValueError(f'{sheet}: orden de días incorrecto.')
                        code = None; profile = None
                        source_range = f'{"A" if offset == 0 else "F"}{day+12}:{"E" if offset == 0 else "J"}{day+12}'
                    if denominator <= 0 or numerator < 0: raise ValueError(f'{sheet}: denominador o numerador inválido.')
                    if metric in PERCENTAGES:
                        numerator = integer(numerator, f'{sheet}/cumple/{day}')
                        if numerator > denominator: raise ValueError(f'{sheet}: cumple excede el total.')
                    result = numerator / denominator * (100 if metric in PERCENTAGES else 1)
                    if not isinstance(cached,(float,int)) or abs(result-cached) > 0.000001:
                        raise ValueError(f'{sheet}: resultado almacenado incoherente en día {day}.')
                    pair[phase] = {'date': stamp.date().isoformat(), 'denominator': denominator, 'numerator': numerator,
                                   'value': result, 'cachedValue': cached, 'sourceRange': source_range,
                                   'code': code, 'profile': profile,
                                   'nature': 'USER_DECLARED_REAL_AGGREGATE' if phase == 'PRETEST' else 'SOURCE_AGGREGATE'}
                if pair['PRETEST']['denominator'] != pair['POSTTEST']['denominator']:
                    raise ValueError(f'{sheet}: denominadores diarios no pareados.')
                if metric == 'TPGR' and pair['PRETEST']['profile'] != pair['POSTTEST']['profile']:
                    raise ValueError('Los perfiles de reporte no coinciden entre fases.')
                rows.append(pair)
            summaries = {}
            for phase, offset in [('PRETEST',0),('POSTTEST',6 if metric == 'TPGR' else 5)]:
                denominator = sum(row[phase]['denominator'] for row in rows)
                numerator = sum(Decimal(str(row[phase]['numerator'])) for row in rows)
                total_row = values[42]
                dcol,ncol,vcol = ((3,4,5) if metric == 'TPGR' else (1,2,3))
                result = float(numerator) / denominator * (100 if metric in PERCENTAGES else 1)
                if denominator != total_row[offset+dcol] or abs(float(numerator)-total_row[offset+ncol]) > 0.000001 or abs(result-total_row[offset+vcol]) > 0.000001:
                    raise ValueError(f'{sheet}: resumen no concilia con los 30 días.')
                if denominator != (30 if metric == 'TPGR' else 1693): raise ValueError(f'{sheet}: total inesperado.')
                summaries[phase] = {'denominator': denominator, 'numerator': float(numerator), 'value': result}
            guides[metric] = {'sheet': sheet, 'title': values[3][1], 'formula': values[6][1], 'definition': values[7][1],
                              'unit': '%' if metric in PERCENTAGES else 's', 'sourceNotice': values[9][0],
                              'note': values[44][0], 'rows': rows, 'summary': summaries}
        for metric in METRICS:
            for i,row in enumerate(guides[metric]['rows']):
                for phase in ['PRETEST','POSTTEST']:
                    base = guides['PRCC']['rows'][i][phase]
                    if row[phase]['date'] != base['date']: raise ValueError('Calendarios no coincidentes entre guías.')
                    if metric != 'TPGR' and row[phase]['denominator'] != base['denominator']:
                        raise ValueError('Distribución diaria incompatible entre guías.')
        discrepancies.extend([
            {'code':'NO_INDIVIDUAL_SOURCE','severity':'LIMITACION', 'message':'El libro contiene agregados diarios sin identificación individual de activos.'},
            {'code':'TPI_PACI_DIFFERENT_BASE','severity':'REVISAR_CRITERIO', 'message':'TPI declara tiempos válidos para 1693 equipos; PACI tiene menos identificaciones correctas. Se conservan ambos denominadores; confirmar el criterio de validez de los tiempos.'},
            {'code':'PAIRING_UNIT','severity':'LIMITACION', 'message':'Los 30 pares diarios están en la fuente. Las transiciones individuales no sirven para inferencias por activo.'},
        ])
        return guides, archive, discrepancies
    finally:
        workbook.close(); formulas.close()


def ranked(codes, salt):
    return sorted(codes, key=lambda code: hashlib.sha256(f'{SEED}:{salt}:{code}'.encode()).hexdigest())


def derive_details(guides, census):
    """Distribución determinista: conteos exactos y tiempos enteros en ms.
    Mantiene el mismo conjunto por día en ambas fases; la identidad se deriva
    porque no existe en el archivo agregado.
    """
    ordered = sorted(census, key=lambda item: item['sample_code'])
    generated = []; cursor = 0
    for day in range(1,31):
        count = guides['PRCC']['rows'][day-1]['PRETEST']['denominator']
        group = ordered[cursor:cursor+count]; cursor += count
        codes = [a['sample_code'] for a in group]
        if len(codes) != count: raise ValueError('El censo no tiene el tamaño requerido por la guía.')
        phase_values = {}
        for phase in ['PRETEST','POSTTEST']:
            selected = {metric: set(ranked(codes,f'{day}:{phase}:{metric}')[:guides[metric]['rows'][day-1][phase]['numerator']]) for metric in PERCENTAGES}
            total_ms = milliseconds(guides['TPI']['rows'][day-1][phase]['numerator'])
            weights = {code: 80 + int(hashlib.sha256(f'{SEED}:{day}:{phase}:time:{code}'.encode()).hexdigest()[:8],16) % 41 for code in codes}
            total_weight = sum(weights.values())
            durations = {code: total_ms*weight//total_weight for code,weight in weights.items()}
            remaining = total_ms - sum(durations.values())
            for code in ranked(codes,f'{day}:{phase}:time-remainder')[:remaining]: durations[code] += 1
            phase_values[phase] = {code: {'day':day, 'date': guides['PRCC']['rows'][day-1][phase]['date'],
                                         'nature':'DERIVED_DISAGGREGATION', 'durationMs':durations[code],
                                         **{metric:code in selected[metric] for metric in PERCENTAGES}} for code in codes}
        for item in group:
            code = item['sample_code']
            generated.append({'sample_code':code, 'snapshot':item['snapshot'],
                              'pre':phase_values['PRETEST'][code], 'post':phase_values['POSTTEST'][code]})
    if cursor != len(census): raise ValueError('Las guías no cubren exactamente el censo.')
    return generated


def reconcile(guides, generated):
    checks = []
    for metric in METRICS:
        for row in guides[metric]['rows']:
            for phase, key in [('PRETEST','pre'),('POSTTEST','post')]:
                source = row[phase]
                if metric == 'TPGR':
                    denominator = source['denominator']; numerator = source['numerator']
                    basis = 'RECALCULO_FUENTE_AGREGADA'
                else:
                    records = [a[key] for a in generated if a[key]['day'] == row['day']]
                    denominator = len(records)
                    numerator = sum(a[metric] for a in records) if metric in PERCENTAGES else sum(a['durationMs'] for a in records)/1000
                    basis = 'DERIVACION_VS_FUENTE'
                checks.append({'metric':metric,'day':row['day'],'phase':phase,'basis':basis,
                               'expectedDenominator':source['denominator'],'actualDenominator':denominator,
                               'expectedNumerator':source['numerator'],'actualNumerator':numerator,
                               'passed':denominator == source['denominator'] and abs(numerator-source['numerator']) < 0.000001
                               and abs(source['value']-(numerator/denominator*(100 if metric in PERCENTAGES else 1))) < 0.000001})
    return checks


def import_guides(path, user_id):
    blob = Path(path).read_bytes(); source_hash = hashlib.sha256(blob).hexdigest()
    guides, archive, findings = parse_guides(blob)
    pairs = db.session.execute(db.select(ResearchSample, Asset).join(Asset)).all()
    census = [{'sample_code':s.sample_code, 'snapshot':{'asset_id':a.id, 'sbn':a.sbn,'description':a.description,'site':a.site,'stratum':s.stratum}} for s,a in pairs]
    census.sort(key=lambda row: row['sample_code'])
    study_id = digest({'source':source_hash,'census':census,'algorithm':ALGORITHM,'seed':SEED,'preNature':'USER_DECLARED_REAL_AGGREGATE'})
    existing = db.session.get(GuideStudy,study_id)
    if existing: return existing, False
    generated = derive_details(guides,census); checks = reconcile(guides,generated)
    if not all(check['passed'] for check in checks): raise ValueError('La derivación no concilia con el archivo.')
    manifest = {'mode':'GUIDE_AUDIT','algorithm':ALGORITHM,'seed':SEED,'population':len(census),
                'sourceHash':source_hash,'guidesHash':digest(guides),'detailsHash':digest(generated),
                'pretestNature':'USER_DECLARED_REAL_AGGREGATE','posttestNature':'SOURCE_AGGREGATE',
                'individualNature':'DERIVED_DISAGGREGATION','findings':findings,
                'checks':len(checks),'passedChecks':sum(check['passed'] for check in checks),
                'statement':'Conciliación aritmética; no certifica autenticidad de mediciones.'}
    try:
        study = GuideStudy(id=study_id,filename=Path(path).name,source_hash=source_hash,source_blob=blob,
                                source_cells=archive,guides=guides,manifest=manifest,created_by=user_id)
        db.session.add(study); db.session.flush()
        for values in generated: db.session.add(GuideAsset(study_id=study_id,**values))
        db.session.add(AuditLog(user_id=user_id,action='IMPORT_GUIDES',entity_type='GUIDE_STUDY',entity_id=study_id,
                       details={'sourceHash':source_hash,'population':len(census),'checks':len(checks),'pretestDeclaration':'Declarado por el usuario; revisar la fuente'}))
        db.session.commit(); return study, True
    except Exception:
        db.session.rollback(); raise
